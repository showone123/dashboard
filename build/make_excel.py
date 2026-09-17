# -*- coding: utf-8 -*-
"""
把 canonical data 导出为「客户上传用」的多 sheet Excel 源数据模板。
数据契约必须与 app 的解析器 / 渲染脚本严格一致。
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BUILD = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BUILD)
OUT = os.path.join(ROOT, "outputs")
os.makedirs(OUT, exist_ok=True)

with open(os.path.join(BUILD, "copper_data.json"), encoding="utf-8") as f:
    D = json.load(f)

HDR_FILL = PatternFill("solid", fgColor="16324F")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="微软雅黑", size=10)
KV_FONT = Font(name="微软雅黑", size=10)
KV_KEY = Font(name="微软雅黑", size=10, bold=True)
THIN = Side(style="thin", color="DDE3EC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FONT = Font(name="微软雅黑", size=12, bold=True, color="16324F")

wb = Workbook()
wb.remove(wb.active)


def style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autofit(ws, ncol, minw=10, maxw=52):
    for c in range(1, ncol + 1):
        w = minw
        for row in ws.iter_rows(min_col=c, max_col=c):
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                s = str(v)
                # 中文字符按 2 个宽度估算
                width = sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)
                w = max(w, width + 3)
        ws.column_dimensions[get_column_letter(c)].width = min(w, maxw)


def add_table(name, headers, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    style_header(ws, len(headers))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.font = CELL_FONT
            cell.border = BORDER
    autofit(ws, len(headers))
    return ws


def add_kv(name, pairs):
    ws = wb.create_sheet(name)
    ws.append(["字段", "值"])
    for k, v in pairs:
        ws.append([k, v])
    style_header(ws, 2)
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2), start=2):
        row[0].font = KV_KEY
        row[1].font = KV_FONT
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
    autofit(ws, 2, minw=18, maxw=60)
    return ws


# ---------- 0. 说明 ----------
ws = wb.create_sheet("说明")
ws["A1"] = "铜期货数据看板 · 源数据模板"
ws["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="16324F")
notes = [
    ("生成时间", D["meta"]["generatedAt"]),
    ("数据来源", D["meta"]["source"]),
    ("数据日期", D["meta"]["date"]),
    ("", ""),
    ("使用方式", "在本文件对应 sheet 中替换数值后另存为 .xlsx，上传到看板即可重新渲染。"),
    ("不要改动", "sheet 名称、表头行、字段名；不要新增/删除 sheet。数值可改，行数可增减（KLINE 最多 400 行）。"),
    ("必填 sheet", "META / SUMMARY / KLINE / CONTRACTS 四张为必填，其余缺失时对应模块显示为空。"),
    ("sheet 清单", "说明、META、SUMMARY、KLINE、TECH、CONTRACTS、ROLL、RANK、RANK_META、WAREHOUSE、INVENTORY、OPTIONS_META、OPTION_CHAIN、OPT_STRATEGY、EXTERNAL、RATIOS、TREND、SENTIMENT、SENTIMENT_WINDOWS、SENTIMENT_EVENTS、SEASON、COST"),
    ("KLINE 顺序", "必须按日期从旧到新排列；日期格式 YYYY-MM-DD 或 MM-DD 均可。"),
    ("价格单位", "元/吨；成交量、持仓量单位为手。"),
    ("", ""),
    ("免责声明", D["meta"]["disclaimer"]),
]
r = 3
for k, v in notes:
    ws.cell(row=r, column=1, value=k).font = KV_KEY
    ws.cell(row=r, column=2, value=v).font = KV_FONT
    ws.cell(row=r, column=2).alignment = Alignment(vertical="center", wrap_text=True)
    r += 1
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 90
for rr in range(3, r):
    ws.row_dimensions[rr].height = 30

# ---------- 1. META ----------
m = D["meta"]
add_kv("META", [
    ("product", m["product"]), ("产品名称", m["productName"]), ("交易所", m["exchange"]),
    ("价格单位", m["unit"]), ("主力合约", m["contract"]), ("连续合约", m["continuous"]),
    ("数据日期", m["date"]), ("生成时间", m["generatedAt"]), ("数据来源", m["source"]),
    ("项目", m["project"]), ("免责声明", m["disclaimer"]),
])

# ---------- 2. SUMMARY ----------
s = D["summary"]
add_kv("SUMMARY", [
    ("最新收盘", s["close"]), ("前收盘", s["prevClose"]), ("涨跌", s["chg"]), ("涨跌幅%", s["chgPct"]),
    ("MA5", s["ma5"]), ("MA20", s["ma20"]), ("ATR14", s["atr"]),
    ("偏离MA5", s["ma5Diff"]), ("偏离MA5%", s["ma5DiffPct"]),
    ("60日最高", s["high60"]), ("60日最低", s["low60"]),
    ("本周最高", s["weekHigh"]), ("本周最低", s["weekLow"]), ("距高点%", s["fromHigh"]),
    ("成交量", s["volume"]), ("持仓量", s["openInterest"]), ("持仓变化", s["oiChange"]),
])

# ---------- 3. KLINE ----------
add_table("KLINE", ["日期", "开盘", "最高", "最低", "收盘", "成交量", "持仓量"],
          [(k["d"], k["o"], k["h"], k["l"], k["c"], k["v"], k["oi"]) for k in D["kline"]])

# ---------- 4. TECH ----------
rows = []
for t in D["tech"]:
    for it in t["items"]:
        rows.append((t["p"], t["price"], it["n"], it["v"], it["d"], t["bias"], "/".join(t["tags"])))
add_table("TECH", ["周期", "周期价格", "指标", "数值", "方向", "综合偏向", "标签"], rows)

# ---------- 5. CONTRACTS ----------
add_table("CONTRACTS", ["合约", "收盘价", "结算价", "成交量", "持仓量", "是否主力"],
          [(c["c"], c["p"], c["set"], c["vol"], c["oi"], "是" if c["c"] == m["contract"] else "")
           for c in D["contracts"]])

# ---------- 6. ROLL ----------
add_table("ROLL", ["组合", "价差", "年化%"],
          [(r["name"], r["spread"], r["annual"]) for r in D["term"]["roll"]])

# ---------- 7. RANK ----------
rows = []
for kind, label in (("volume", "成交量"), ("long", "多头持仓"), ("short", "空头持仓")):
    for it in D["rank"][kind]:
        rows.append((label, it["rank"], it["name"], it["value"], it["change"]))
add_table("RANK", ["类型", "排名", "会员", "数值", "变化"], rows)

rk = D["rank"]["top20"]
add_kv("RANK_META", [
    ("前20成交量", rk["volume"]), ("前20成交量变化", rk["volumeChange"]),
    ("前20多头", rk["long"]), ("前20多头变化", rk["longChange"]),
    ("前20空头", rk["short"]), ("前20空头变化", rk["shortChange"]),
    ("净多净空差", D["rank"]["netLongShort"]), ("信号", D["rank"]["signal"]),
])

# ---------- 8. WAREHOUSE ----------
w = D["warehouse"]
add_kv("WAREHOUSE", [
    ("最新仓单", w["stocks"]), ("变化", w["change"]), ("交易所", w["exchange"]),
    ("日期", w["date"]), ("解读", w["note"]),
])

# ---------- 9. INVENTORY ----------
add_table("INVENTORY", ["指标", "数值", "单位", "截止日期", "备注"],
          [(i["name"], i["value"], i["unit"], i["asOf"], i["note"]) for i in D["inventory"]])

# ---------- 10. OPTIONS_META ----------
o = D["options"]
add_kv("OPTIONS_META", [
    ("标的合约", o["contract"]), ("数据日期", o["date"]), ("ATM行权价", o["atmStrike"]),
    ("Call价格", o["callPrice"]), ("Put价格", o["putPrice"]), ("CallDelta", o["callDelta"]),
    ("标的期货价", o["futuresPrice"]), ("MaxPain", o["maxPain"]),
    ("HV20", o["hv20"]), ("HV60", o["hv60"]),
    ("PCR(OI)", o["pcrOi"]), ("PCR(Vol)", o["pcrVolume"]),
    ("总CallOI", o["totalCallOi"]), ("总PutOI", o["totalPutOi"]),
    ("PCR信号", o["pcrSignal"]), ("波动率解读", o["volNote"]),
])

# ---------- 11. OPTION_CHAIN ----------
add_table("OPTION_CHAIN",
          ["行权价", "Call结算", "Call量", "Call持仓", "CallDelta", "Put结算", "Put量", "Put持仓", "PutDelta", "是否ATM"],
          [(c["k"], c["cp"], c["cv"], c["co"], c["cd"], c["pp"], c["pv"], c["po"], c["pd"],
            "是" if c.get("atm") else "") for c in o["chain"]])

# ---------- 12. OPT_STRATEGY ----------
rows = []
for st in o["strategies"]:
    for k, v in st["rows"]:
        rows.append((st["name"], k, v))
add_table("OPT_STRATEGY", ["策略", "项目", "内容"], rows)

# ---------- 13. EXTERNAL ----------
e = D["external"]
add_kv("EXTERNAL", [
    ("品种", e["name"]), ("代码", e["symbol"]), ("最新价", e["close"]), ("日期", e["date"]),
    ("1日涨跌%", e["chg1d"]), ("5日涨跌%", e["chg5d"]), ("MA5", e["ma5"]), ("MA20", e["ma20"]),
    ("计价单位", "美元/磅"),
])

# ---------- 14. RATIOS ----------
add_table("RATIOS", ["比价", "当前值", "10年分位", "信号", "方向"],
          [(r["n"], r["v"], r["pct"], r["sig"], r["d"]) for r in D["ratios"]])

# ---------- 15. TREND ----------
cuag, auag = D["trend"]["cuag"], D["trend"]["auag"]
n = max(len(cuag), len(auag))
add_table("TREND", ["序号", "铜银比", "金银比"],
          [(i + 1, cuag[i] if i < len(cuag) else "", auag[i] if i < len(auag) else "") for i in range(n)])

# ---------- 16. SENTIMENT ----------
sn = D["sentiment"]
add_kv("SENTIMENT", [
    ("快照时间", sn["snapshot"]), ("48h条数", sn["total48h"]), ("24h条数", sn["total24h"]),
    ("高影响力事件", sn["alertCount"]), ("当前信号", sn["signal"]), ("动量方向", sn["momentum"]),
    ("边际变化", sn["shift"]), ("中长期", sn["weekly"]), ("宏观占比说明", sn["macroNote"]),
])
add_table("SENTIMENT_WINDOWS", ["窗口", "打分", "判定", "条数"],
          [(wd["label"], wd["score"], wd["verdict"], wd["count"]) for wd in sn["windows"]])

# ---------- 17. SENTIMENT_EVENTS ----------
add_table("SENTIMENT_EVENTS", ["评分", "标签", "时间", "内容"],
          [(ev["sc"], ev["tag"], ev["time"], ev["txt"]) for ev in sn["events"]])

# ---------- 18. SEASON ----------
add_table("SEASON", ["月份", "平均收益率%", "胜率%"],
          [(s["m"], s["r"], s["w"]) for s in D["season"]])

# ---------- 19. COST ----------
c = D["cost"]
add_kv("COST", [
    ("计算价格", c["price"]), ("合约乘数", c["volumeMultiple"]), ("最小变动价位", c["priceTick"]),
    ("开仓手续费", c["openFee"]), ("平仓手续费", c["closeFee"]), ("平今手续费", c["closeTodayFee"]),
    ("开平合计", c["roundTrip"]),
    ("交易所开仓", c["exOpen"]), ("交易所平仓", c["exClose"]), ("交易所开平合计", c["exRoundTrip"]),
    ("公司单边费率%", c["commissionRate"]), ("交易所费率%", c["exRate"]),
    ("公司平今费率%", c["closeTodayRate"]),
    ("保证金率%", c["marginRate"]), ("1手保证金", c["marginPerLot"]), ("保证金模型", c["marginModel"]),
])

path = os.path.join(OUT, "铜期货源数据_20260915.xlsx")
wb.save(path)
print("saved:", path)
print("sheets:", len(wb.sheetnames))
print(wb.sheetnames)
